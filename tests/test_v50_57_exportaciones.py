"""v50.57 · Revisión general (15/10/2026), Bloque 1.3: exportaciones.

- Todas las exportaciones de la web (centro, freelance y administración) descargan en Excel (.xlsx, por defecto) y en
  CSV (UTF-8 con BOM y «;»), con datos y sin datos (vacías: las cabeceras y «Sin datos en este periodo»).
- El Excel lleva las columnas con nombre en español, la primera fila fija y con filtro, fechas y euros con formato.
- El periodo (mes, trimestre, año, temporada, todo o dos fechas) filtra de verdad.
- DAC7: ejercicio fiscal, identificación fiscal, país, importes y actividades por trimestre y comisiones (0); con un
  centro sin datos fiscales el informe sale igualmente y la pestaña «Datos incompletos» lo lista.
- Cada uno ve solo lo suyo: un buceador no descarga nada, un centro no descarga lo de la administración, un perfil del
  equipo sin «reservas» no descarga reservas, y los enlaces de antes (/api/center/export, /api/admin/dac7) siguen.
"""
from __future__ import annotations

import datetime as dt
import io
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from helpers import IN_PERSON, ServerTestCase  # noqa: E402

try:
    import openpyxl
except ImportError:   # the Excel files are then checked by their XML
    openpyxl = None

TODAY = dt.date.today()
DAY = (TODAY + dt.timedelta(days=3)).isoformat()
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PROVIDER_KINDS = ('reservas', 'clientes', 'servicios', 'pagos', 'agenda', 'actividad', 'analitica', 'ventas_canal')
ADMIN_KINDS = ('actividad', 'pagos', 'reservas_centros', 'cuotas')


def sheets_of(raw: bytes) -> dict:
    """{tab: [rows]} of an Excel file, read as Excel would (openpyxl) or, without it, at least its structure."""
    if openpyxl:
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        assert '[Content_Types].xml' in z.namelist() and 'xl/workbook.xml' in z.namelist()
        return {n: [] for n in z.namelist() if n.startswith('xl/worksheets/')}


class Exportaciones(ServerTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        t = cls('run')
        cls.admin = t.make_admin()
        cls.center = t.make_center(cls.admin, 'expo')
        cls.pro = t.make_freelance(cls.admin, 'expo')
        cls.diver = t.make_diver('expo')
        cid = cls.center.user['center_id']
        sid = cls.center.post('/api/provider-services', {'title': 'Bautizo para exportar', 'price': 45.5, 'booking_mode': 'request', 'max_participants': 6}).json['id']
        slot = cls.center.post('/api/provider-services/slot', {'service_id': sid, 'slot_date': DAY, 'slot_time': '10:00', 'spots_total': 6}).json['slot_id']
        r = cls.diver.post('/api/bookings', {'center_id': cid, 'service_id': sid, 'slot_id': slot, 'participants': 2})
        assert r.status == 200, r.json
        cls.bid = r.json['id']
        cls.center.post(f'/api/bookings/{cls.bid}/status', {'status': 'confirmed'})
        psid = cls.pro.post('/api/provider-services', {'title': 'Guía para exportar', 'price': 60, 'booking_mode': 'request', 'max_participants': 4, **IN_PERSON}).json['id']
        pslot = cls.pro.post('/api/provider-services/slot', {'service_id': psid, 'slot_date': DAY, 'slot_time': '09:00', 'spots_total': 4}).json['slot_id']
        cls.diver.post('/api/bookings', {'service_id': psid, 'slot_id': pslot, 'participants': 1, 'in_person_ack': True})
        # a payment through the platform in a past year, for a centre without fiscal data (DAC7)
        cls.year = TODAY.year - 1
        cls.server.sql("""INSERT INTO payment_orders(booking_id,user_id,service_id,amount,currency,status,created_at,updated_at,paid_at,payee_type,payee_id,payment_kind)
                          VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""", (cls.bid, cls.diver.user['id'], sid, 91.0, 'EUR', 'paid', f'{cls.year}-05-10T10:00:00', f'{cls.year}-05-10T10:01:00',
                                                                f'{cls.year}-05-10T10:01:00', 'center', cid, 'full'))

    def get(self, who, kind, **q):
        q.setdefault('formato', 'xlsx'); q.setdefault('periodo', 'todo')
        return who.get(f'/api/exportar/{kind}?' + '&'.join(f'{k}={v}' for k, v in q.items()))

    def assert_csv(self, r, header_start):
        self.assertEqual(r.status, 200, r.json)
        self.assertIn('text/csv', r.headers.get('Content-Type', ''))
        self.assertIn('attachment', r.headers.get('Content-Disposition', '')); self.assertIn('.csv"', r.headers.get('Content-Disposition', ''))
        self.assertTrue(r.body.startswith(b'\xef\xbb\xbf'), 'sin BOM')
        lines = r.body.decode('utf-8-sig').split('\r\n')
        self.assertTrue(lines[0].startswith(header_start), lines[0]); self.assertIn(';', lines[0])
        return lines

    def assert_xlsx(self, r):
        self.assertEqual(r.status, 200, r.json)
        self.assertEqual(r.headers.get('Content-Type'), XLSX); self.assertIn('.xlsx"', r.headers.get('Content-Disposition', ''))
        return sheets_of(r.body)

    def test_centro_todas_en_excel_y_csv_con_datos(self):
        for kind in PROVIDER_KINDS:
            with self.subTest(kind):
                tabs = self.assert_xlsx(self.get(self.center, kind))
                self.assertTrue(tabs)
                self.assert_csv(self.get(self.center, kind, formato='csv'), '')
        tabs = self.assert_xlsx(self.get(self.center, 'reservas'))
        rows = tabs['Reservas'] if openpyxl else None
        if rows:
            self.assertEqual(rows[0][:4], ['Reserva', 'Fecha', 'Hora', 'Actividad'])
            row = next(x for x in rows[1:] if x[0] == f'WOD-{self.bid:06d}')
            price = self.server.sql('SELECT price FROM bookings WHERE id=?', (self.bid,))[0]['price']
            self.assertIsInstance(row[1], dt.datetime); self.assertEqual(row[5], 2); self.assertEqual(row[8], round(price, 2)); self.assertEqual(row[9], 'EUR')
        lines = self.assert_csv(self.get(self.center, 'reservas', formato='csv'), 'Reserva;Fecha;Hora;Actividad')
        line = next(x for x in lines if x.startswith(f'WOD-{self.bid:06d}'))
        self.assertIn(dt.date.fromisoformat(DAY).strftime('%d/%m/%Y'), line)    # the date as Excel in Spanish writes it

    @unittest.skipUnless(openpyxl, 'openpyxl no está instalado')
    def test_excel_profesional(self):
        r = self.get(self.center, 'reservas')
        ws = openpyxl.load_workbook(io.BytesIO(r.body))['Reservas']
        self.assertEqual(ws.freeze_panes, 'A2'); self.assertTrue(ws.auto_filter.ref.startswith('A1:'))
        self.assertTrue(ws['A1'].font.b)
        fmts = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c).number_format for c in range(1, ws.max_column + 1)}
        self.assertEqual(fmts['Fecha'], 'dd/mm/yyyy'); self.assertIn('€', fmts['Importe']); self.assertEqual(fmts['Creada'], 'dd/mm/yyyy hh:mm')

    def test_vacias_descargan_con_cabeceras_y_aviso(self):
        far = {'periodo': 'personalizado', 'desde': '2001-01-01', 'hasta': '2001-01-31'}
        for kind in PROVIDER_KINDS:
            with self.subTest(kind):
                lines = self.assert_csv(self.get(self.center, kind, formato='csv', **far), '')
                if kind != 'servicios': self.assertTrue(lines[1].startswith('Sin datos en este periodo'), lines[:3])
                tabs = self.assert_xlsx(self.get(self.center, kind, **far))
                if openpyxl and kind != 'servicios': self.assertEqual(list(tabs.values())[0][1][0], 'Sin datos en este periodo')

    def test_el_periodo_filtra(self):
        month = self.assert_csv(self.get(self.center, 'reservas', formato='csv', periodo='mes'), 'Reserva')
        in_month = DAY[:7] == TODAY.isoformat()[:7]
        self.assertEqual(any(x.startswith(f'WOD-{self.bid:06d}') for x in month), in_month)
        rng = {'periodo': 'personalizado', 'desde': DAY, 'hasta': DAY}
        self.assertTrue(any(x.startswith(f'WOD-{self.bid:06d}') for x in self.assert_csv(self.get(self.center, 'reservas', formato='csv', **rng), 'Reserva')))
        for per in ('trimestre', 'trimestre_anterior', 'anio', 'anio_anterior', 'temporada', 'mes_anterior'):
            with self.subTest(per): self.assertEqual(self.get(self.center, 'reservas', formato='csv', periodo=per).status, 200)
        r = self.get(self.center, 'reservas', periodo='personalizado', desde='', hasta='')
        self.assertEqual(r.status, 400); self.assertIn('fechas', r.json['error'])

    def test_freelance(self):
        for kind in ('reservas', 'clientes', 'servicios', 'pagos', 'agenda', 'ventas_canal'):
            with self.subTest(kind): self.assert_xlsx(self.get(self.pro, kind))
        lines = self.assert_csv(self.get(self.pro, 'servicios', formato='csv'), 'Id;Servicio')
        self.assertTrue(any('Guía para exportar' in x for x in lines))
        self.assertFalse(any('Bautizo para exportar' in x for x in lines))     # never another provider's

    def test_administracion_y_dac7(self):
        for kind in ADMIN_KINDS:
            with self.subTest(kind):
                self.assert_xlsx(self.get(self.admin, kind)); self.assert_csv(self.get(self.admin, kind, formato='csv'), '')
        r = self.admin.get(f'/api/exportar/dac7?formato=xlsx&anio={self.year}')
        tabs = self.assert_xlsx(r)
        if openpyxl:
            main = tabs[f'DAC7 {self.year}']
            self.assertEqual(main[0][:7], ['Ejercicio', 'Tipo de prestador', 'Id', 'Nombre', 'Razón social', 'Persona', 'NIF / TIN'])
            row = next(x for x in main[1:] if x[2] == self.center.user['center_id'])
            head = main[0]
            self.assertEqual(row[0], self.year); self.assertEqual(row[head.index('Importe T2')], 91.0); self.assertEqual(row[head.index('Actividades T2')], 1)
            self.assertEqual(row[head.index('Comisiones')], 0); self.assertIn('País emisor', head)
            self.assertTrue(row[head.index('Datos incompletos')])                                        # no fiscal data
            self.assertTrue(any(x[1] == self.center.user['center_id'] for x in tabs['Datos incompletos'][1:]))
        lines = self.assert_csv(self.admin.get(f'/api/exportar/dac7?formato=csv&anio={self.year}'), 'Ejercicio;Tipo de prestador')
        self.assertTrue(any(';91,00;' in x for x in lines), lines[:3])
        # an empty year downloads too
        self.assert_csv(self.admin.get('/api/exportar/dac7?formato=csv&anio=2021'), 'Ejercicio')
        self.assert_xlsx(self.admin.get('/api/exportar/dac7_pago_en_persona?formato=xlsx&anio=2021'))

    def test_enlaces_de_antes_siguen(self):
        self.assert_csv(self.center.get('/api/center/export?kind=bookings'), 'Reserva;Fecha')
        self.assert_csv(self.center.get('/api/center/export?kind=clients'), 'Nombre;Email')
        self.assert_csv(self.admin.get(f'/api/admin/dac7?format=csv&year={self.year}'), 'Ejercicio')
        self.assert_xlsx(self.admin.get(f'/api/admin/dac7?format=xlsx&year={self.year}'))

    def test_cada_uno_lo_suyo(self):
        self.assertEqual(self.get(self.diver, 'reservas').status, 403)
        self.assertIn(self.client().get('/api/exportar/reservas').status, (401, 403))   # without a session
        for kind in ('dac7', 'cuotas', 'reservas_centros'):
            with self.subTest(kind): self.assertEqual(self.get(self.center, kind).status, 403)
        for kind in ('reservas', 'clientes', 'agenda'):
            with self.subTest(kind): self.assertEqual(self.get(self.admin, kind).status, 403)
        self.assertEqual(self.get(self.center, 'inventada').status, 404)
        # the downloads are written in the activity
        self.assertTrue(self.server.sql("SELECT 1 FROM audit_log WHERE action IN ('provider_export','data_exported','dac7_exported')"))


if __name__ == '__main__':
    unittest.main()
